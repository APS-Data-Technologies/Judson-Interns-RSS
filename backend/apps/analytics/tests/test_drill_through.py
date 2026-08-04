from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APITestCase
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.leads.models import Family, LeadSource
from apps.sites.models import Location
from apps.reports.models import CostBasis
from apps.tours.models import Tour, TourEvent, TourStatus


def aware(year, month, day):
    return timezone.make_aware(datetime(year, month, day, 12))


class AnalyticsDrillThroughTests(APITestCase):
    def setUp(self):
        self.location = Location.objects.create(
            location_name="Downtown",
            address="100 Main Street",
            city="Chicago",
            state="IL",
            zip_code="60601",
        )
        self.staff = User.objects.create_user(
            email="staff@example.com",
            password="test-pass",
            first_name="Staff",
            last_name="User",
            role=User.Role.STAFF,
            location=self.location,
        )
        self.source = LeadSource.objects.create(source_name="Referral")
        self.url = reverse("analytics-drill-through")

        self.current_booked = self.create_tour("Current Booked", aware(2026, 7, 15))
        self.current_toured = self.create_tour("Current Toured", aware(2026, 7, 16), TourStatus.TOURED)
        self.previous_booked = self.create_tour("Previous Booked", aware(2026, 6, 15))

        Tour.objects.filter(pk=self.current_booked.pk).update(created_at=aware(2026, 7, 10))
        Tour.objects.filter(pk=self.current_toured.pk).update(created_at=aware(2026, 7, 11))
        Tour.objects.filter(pk=self.previous_booked.pk).update(created_at=aware(2026, 6, 10))
        TourEvent.objects.create(
            tour=self.current_toured,
            status=TourStatus.TOURED,
            event_timestamp=aware(2026, 7, 18),
            updated_by=self.staff,
        )

        self.client.force_authenticate(self.staff)

    def create_tour(self, family_name, scheduled_date, status=TourStatus.SCHEDULED):
        return Tour.objects.create(
            family=Family.objects.create(
                family_name=family_name,
                contact_email="family@example.com",
                contact_phone="210-555-0100",
            ),
            location=self.location,
            lead_source=self.source,
            assigned_staff=self.staff,
            student_name="Alex",
            scheduled_tour_date=scheduled_date,
            current_status=status,
        )

    def test_event_drill_through_uses_booked_creation_date(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_mode": "event",
            "drill_status": "scheduled",
            "drill_kpi": "Booked",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 2)
        self.assertEqual(
            {row["familyName"] for row in response.data["rows"]},
            {"Current Booked", "Current Toured"},
        )
        self.assertTrue(all(row["contributionDate"].startswith("2026-07") for row in response.data["rows"]))
        self.assertTrue(all(row["contributingKpi"] == "Booked" for row in response.data["rows"]))
        self.assertTrue(all(row["studentName"] == "Alex" for row in response.data["rows"]))
        self.assertTrue(all(row["emailPhone"] == "family@example.com · 210-555-0100" for row in response.data["rows"]))

    def test_event_drill_through_uses_selected_status_event_date(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_mode": "event",
            "drill_status": "toured",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["rows"][0]["familyName"], "Current Toured")
        self.assertEqual(response.data["rows"][0]["contributionDate"], "2026-07-18")

    def test_cohort_drill_through_uses_scheduled_period_and_reached_status(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_mode": "cohort",
            "drill_status": "toured",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["rows"][0]["familyName"], "Current Toured")
        self.assertEqual(response.data["rows"][0]["contributionDate"], "2026-07-16")

    def test_rate_drill_through_returns_denominator_and_marks_numerator(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_mode": "cohort",
            "drill_denominator": "scheduled",
            "drill_numerator": "toured",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 2)
        contributions = {
            row["familyName"]: row["metricRole"]
            for row in response.data["rows"]
        }
        self.assertEqual(contributions["Current Toured"], "Numerator")
        self.assertEqual(contributions["Current Booked"], "Denominator only")

    def test_chart_drill_through_combines_all_requested_statuses(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_statuses": "scheduled,toured",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 3)
        self.assertEqual(
            [row["contributingKpi"] for row in response.data["rows"]].count("Booked"),
            2,
        )
        self.assertEqual(
            [row["contributingKpi"] for row in response.data["rows"]].count("Toured"),
            1,
        )

    def test_dimension_drill_through_limits_rows_to_selected_entity(self):
        other_location = Location.objects.create(
            location_name="North",
            address="200 North Street",
            city="Chicago",
            state="IL",
            zip_code="60602",
        )
        other_tour = self.create_tour("Other Location", aware(2026, 7, 17))
        Tour.objects.filter(pk=other_tour.pk).update(
            created_at=aware(2026, 7, 12),
            location=other_location,
        )

        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_mode": "event",
            "drill_status": "scheduled",
            "drill_dimension": "location",
            "drill_dimension_value": "Downtown",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 2)
        self.assertTrue(all(row["location"] == "Downtown" for row in response.data["rows"]))

    def test_temporal_drill_through_limits_rows_to_selected_bucket(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_mode": "event",
            "drill_status": "scheduled",
            "drill_temporal_dimension": "dayOfMonth",
            "drill_temporal_value": "10",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["rows"][0]["familyName"], "Current Booked")

    def test_all_time_scope_ignores_selected_page_period(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_mode": "event",
            "drill_status": "scheduled",
            "drill_scope": "all",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 3)
        self.assertIn("Previous Booked", {row["familyName"] for row in response.data["rows"]})

    def test_time_to_progress_drill_through_uses_transition_and_elapsed_bucket(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_transition": "booked_to_toured",
            "drill_elapsed_min": "0",
            "drill_elapsed_max": "7",
            "drill_kpi": "Booked to Toured · 0–7 days",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["rows"][0]["familyName"], "Current Toured")
        self.assertEqual(response.data["rows"][0]["elapsedDays"], 7)
        self.assertEqual(response.data["rows"][0]["elapsedBucket"], "0–7 days")

    def test_time_to_progress_section_drill_combines_requested_transitions(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_transitions": "booked_to_toured,toured_to_enrolled",
            "drill_kpi": "Time to Progress",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["rows"][0]["contributingKpi"], "Booked → Toured")
        self.assertEqual(response.data["rows"][0]["elapsedDays"], 7)
        self.assertEqual(response.data["rows"][0]["elapsedBucket"], "0–7 days")

    def test_pending_tour_outcome_drill_through_returns_only_overdue_unresolved_tours(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_pending": "tour_outcome",
            "drill_kpi": "Pending Toured / No Show",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["rows"][0]["familyName"], "Current Booked")

    def test_financial_drill_through_returns_cost_basis_entries(self):
        admin = User.objects.create_user(
            email="admin@example.com",
            password="test-pass",
            role=User.Role.ADMIN,
        )
        self.client.force_authenticate(admin)
        CostBasis.objects.create(
            location=self.location,
            reporting_month=aware(2026, 7, 1).date(),
            cost_type=CostBasis.CostType.REVENUE,
            cost_amount=Decimal("12500.00"),
        )
        CostBasis.objects.create(
            location=self.location,
            reporting_month=aware(2026, 7, 1).date(),
            cost_type=CostBasis.CostType.EXPENDITURE,
            cost_amount=Decimal("4500.00"),
        )

        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_financial": "revenue",
            "drill_kpi": "Revenue",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["rows"][0]["type"], "Revenue")
        self.assertEqual(response.data["rows"][0]["amount"], "12500.00")

    def test_drill_through_export_contains_exact_rows_and_context(self):
        response = self.client.get(self.url, {
            "date_from": "2026-07-01",
            "date_to": "2026-07-31",
            "drill_mode": "event",
            "drill_status": "toured",
            "drill_kpi": "Toured",
            "drill_page": "Overview",
            "drill_section": "Volume and Trend Analysis",
            "drill_visualization": "Selected period volume",
            "drill_current_value": "1",
            "drill_previous_value": "3",
            "drill_difference": "-2",
            "drill_difference_percent": "-67%",
            "drill_current_period": "July 1, 2026 to July 31, 2026",
            "drill_previous_period": "June 1, 2026 to June 30, 2026",
            "drill_location_label": "All locations",
            "drill_lead_source_label": "All lead sources",
            "drill_staff_label": "All staff",
            "export": "xlsx",
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Disposition"],
            'attachment; filename="Overview - Volume and Trend Analysis - Selected period volume - Toured.xlsx"',
        )
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook["Drill Through"]
        values = list(sheet.iter_rows(values_only=True))
        data_title_index = next(index for index, row in enumerate(values) if row[0] == "Contributing Data")
        context = dict(row[:2] for row in values[1:data_title_index - 1])
        rows = values[data_title_index + 1:]
        self.assertEqual(context["Page"], "Overview")
        self.assertEqual(context["Section"], "Volume and Trend Analysis")
        self.assertEqual(context["Chart / card"], "Selected period volume")
        self.assertEqual(context["Previous Period"], "June 1, 2026 to June 30, 2026")
        self.assertEqual(context["Location"], "All locations")
        self.assertEqual(context["Lead Source"], "All lead sources")
        self.assertEqual(context["Staff"], "All staff")
        self.assertEqual(context["Difference %"], "-67%")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], "Toured")
        self.assertEqual(rows[1][2], "Current Toured")
