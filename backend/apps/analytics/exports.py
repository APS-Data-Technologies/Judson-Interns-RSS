from io import BytesIO
from datetime import datetime
import re
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .services import cohort_analytics


PAGE_LABELS = {
    "overview": "Overview",
    "volume": "Volume & Trend",
    "cohort": "Conversion & Cohort",
    "locations": "Location",
    "leadSources": "Lead Source",
    "staff": "Staff",
    "cost-margin": "Costs & Margin",
}
PUBLIC_PAGES = ["overview", "volume", "cohort", "locations", "leadSources"]
RESTRICTED_PAGES = ["staff", "cost-margin"]
NAVY = colors.HexColor("#163968")
NAVY_DARK = colors.HexColor("#071B40")
GOLD = colors.HexColor("#FFBF21")
BLUE = colors.HexColor("#0878DE")
SOFT_BLUE = colors.HexColor("#EEF5FC")
MUTED = colors.HexColor("#65758D")
GRID = colors.HexColor("#D7E1ED")


def allowed_pages(user):
    role = getattr(user, "role", "")
    return PUBLIC_PAGES + (RESTRICTED_PAGES if role in {"admin", "super_admin"} else [])


def _rows(data, page):
    if page == "overview":
        labels = {
            "scheduled": "Booked",
            "toured": "Toured",
            "no_show": "No Show",
            "enrolled": "Enrolled",
            "churned": "Churned",
        }
        return [
            ["Metric", "Current", "Previous"],
            *[
                [labels.get(key, key), value, data.get("previousCounts", {}).get(key, 0)]
                for key, value in data.get("counts", {}).items()
            ],
        ]
    if page == "volume":
        return [
            ["Period", "Booked", "Toured", "No Show", "Enrolled", "Churned"],
            *[
                [
                    row.get("label"),
                    row.get("scheduled", 0),
                    row.get("toured", 0),
                    row.get("no_show", 0),
                    row.get("enrolled", 0),
                    row.get("churned", 0),
                ]
                for row in data.get("volumeTrendData", [])
            ],
        ]
    if page == "cohort":
        rates = data.get("rates", {})
        return [
            ["Metric", "Value", "Change"],
            ["Toured Rate", rates.get("toured"), data.get("rateDeltas", {}).get("toured")],
            ["No Show Rate", rates.get("no_show"), data.get("rateDeltas", {}).get("no_show")],
            ["Close Rate", rates.get("close"), data.get("rateDeltas", {}).get("close")],
            ["Conversion Rate", rates.get("conversion"), data.get("rateDeltas", {}).get("conversion")],
            ["Average Days to Enroll", data.get("averageDaysToEnroll"), ""],
        ]
    if page in {"locations", "leadSources", "staff"}:
        return _entity_rows(data.get("entityHealth", {}).get(page, []))
    if page == "cost-margin":
        financial = data.get("financialSummary", {})
        return [
            ["Metric", "Value"],
            ["Revenue", financial.get("revenue", 0)],
            ["Costs", financial.get("cost", 0)],
            ["Contribution Margin", financial.get("margin", 0)],
            [],
            ["Month", "Revenue", "Costs", "Contribution Margin", "Margin %"],
            *[
                [row.get("label"), row.get("revenue"), row.get("cost"), row.get("margin"), row.get("marginRate")]
                for row in data.get("financialTrend", [])
            ],
        ]
    return [["No data"]]


def _entity_rows(rows):
    headers = ["Name", "Booked", "Toured", "No Show", "Enrolled", "Churned", "Conversion %"]
    output = [headers]
    for row in rows:
        current = row.get("current", row)
        output.append([
            row.get("name", row.get("label", "—")),
            current.get("scheduled", current.get("booked", 0)),
            current.get("toured", 0),
            current.get("no_show", current.get("noShow", 0)),
            current.get("enrolled", 0),
            current.get("churned", 0),
            current.get("conversionRate", current.get("conversion", 0)),
        ])
    return output


def export_analytics(user, payload):
    coverage = payload.get("coverage", "current")
    current_page = payload.get("page", "overview")
    pages = allowed_pages(user)
    selected_pages = pages if coverage == "all" else [current_page]
    selected_pages = [page for page in selected_pages if page in pages]

    params = dict(payload.get("filters") or {})
    if payload.get("viewScope") == "all":
        for key in ("metric", "ranking_metric", "ranking_sort", "cost_basis"):
            params.pop(key, None)
    if payload.get("dataScope") == "all_authorized":
        for key in ("location", "locations", "lead_source", "lead_sources", "staff"):
            params.pop(key, None)
    data = cohort_analytics(user, params)
    if payload.get("format") == "pdf":
        return (
            _pdf(data, selected_pages, params),
            "application/pdf",
            _pdf_filename(payload, selected_pages),
        )
    return _excel(data, selected_pages, params), (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ), "analytics-data.xlsx"


def _pdf_filename(payload, pages):
    coverage_label = "All Pages" if payload.get("coverage") == "all" else PAGE_LABELS.get(
        pages[0] if pages else payload.get("page", "overview"),
        "Analytics",
    )
    parts = ["RSS Analytics", coverage_label]
    parts.append("All Views" if payload.get("viewScope") == "all" else "Current View")
    parts.append("All Authorized Data" if payload.get("dataScope") == "all_authorized" else "Current Filtered Data")

    filename = " - ".join(_safe_filename_part(part) for part in parts if part)
    return f"{filename[:180].rstrip(' .-')}.pdf"


def _safe_filename_part(value):
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "-", str(value))
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .-")
    return cleaned or "Analytics"


def export_drill_through(rows, params):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Drill Through"
    sheet.append(["Context", "Value"])
    context_rows = [
        ("KPI", params.get("drill_kpi")),
        ("Current Value", params.get("drill_current_value")),
        ("Previous Value", params.get("drill_previous_value")),
        ("Difference", params.get("drill_difference")),
        ("Difference %", params.get("drill_difference_percent")),
        ("Page", params.get("drill_page")),
        ("Section", params.get("drill_section")),
        ("Chart / card", params.get("drill_visualization")),
        ("Scope", "All time" if params.get("drill_scope") == "all" else "Selected period"),
        ("Current Period", params.get("drill_current_period")),
        ("Previous Period", params.get("drill_previous_period")),
        ("Location", params.get("drill_location_label")),
        ("Lead Source", params.get("drill_lead_source_label")),
        ("Staff", params.get("drill_staff_label")),
    ]
    for label, value in context_rows:
        if value not in (None, ""):
            sheet.append([label, value])

    sheet.append([])
    sheet.append(["Contributing Data"])
    tour_columns = [
        ("contributingKpi", "Contributing KPI"),
        ("contributionDate", "Contribution Date"),
        *(
            [
                ("elapsedDays", "Elapsed Days"),
                ("elapsedBucket", "Days Bucket"),
            ]
            if any(row.get("elapsedDays") is not None for row in rows)
            else []
        ),
        ("familyName", "Family Name"),
        ("scheduledDateTime", "Scheduled Tour"),
        ("location", "Location"),
        ("currentStatus", "Current Status"),
        ("assignedStaff", "Assigned Staff"),
        ("leadSource", "Lead Source"),
        ("studentName", "Student Name"),
        ("childGrade", "Child Grade"),
        ("emailPhone", "Email / Phone"),
    ]
    financial_columns = [
        ("contributingKpi", "Contributing KPI"),
        ("reportingMonth", "Reporting Month"),
        ("location", "Location"),
        ("type", "Type"),
        ("amount", "Amount"),
        ("notes", "Notes"),
    ]
    columns = tour_columns if any(row.get("tourId") for row in rows) else financial_columns
    data_header_row = sheet.max_row + 1
    sheet.append([label for _, label in columns])
    for row in rows:
        sheet.append([row.get(key, "—") for key, _ in columns])

    sheet.freeze_panes = f"A{data_header_row + 1}"
    sheet.auto_filter.ref = f"A{data_header_row}:{sheet.cell(sheet.max_row, len(columns)).coordinate}"
    for header_row in (1, data_header_row):
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163968")
    sheet.cell(data_header_row - 1, 1).font = Font(bold=True, color="163968", size=14)
    for column in sheet.columns:
        width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _excel(data, pages, params):
    workbook = Workbook()
    workbook.remove(workbook.active)
    filters = workbook.create_sheet("Filters & Definitions")
    filters.append(["Filter", "Value"])
    for key, value in sorted(params.items()):
        filters.append([key.replace("_", " ").title(), ", ".join(value) if isinstance(value, list) else value])
    filters.append([])
    filters.append(["Definition", "Meaning"])
    filters.append(["Current selections", "Values and page controls selected when the export was generated."])
    filters.append(["All authorized data", "Entity filters removed; role and location permissions remain enforced."])

    for page in pages:
        sheet = workbook.create_sheet(PAGE_LABELS[page][:31])
        for row in _rows(data, page):
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163968")
        for column in sheet.columns:
            width = min(40, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _pdf(data, pages, params):
    stream = BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=landscape(letter),
        leftMargin=0.52 * inch,
        rightMargin=0.52 * inch,
        topMargin=0.78 * inch,
        bottomMargin=0.55 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "AnalyticsTitle",
        parent=styles["Title"],
        textColor=GOLD,
        fontName="Helvetica-Bold",
        fontSize=28,
        leading=32,
        alignment=TA_CENTER,
        spaceAfter=10,
    )
    cover_subtitle = ParagraphStyle(
        "AnalyticsCoverSubtitle",
        parent=styles["BodyText"],
        textColor=colors.white,
        fontSize=11,
        leading=16,
        alignment=TA_CENTER,
    )
    page_title = ParagraphStyle(
        "AnalyticsPageTitle",
        parent=styles["Heading1"],
        textColor=NAVY_DARK,
        fontName="Helvetica-Bold",
        fontSize=19,
        leading=23,
        spaceAfter=5,
    )
    section_label = ParagraphStyle(
        "AnalyticsSectionLabel",
        parent=styles["BodyText"],
        textColor=BLUE,
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        spaceBefore=4,
        spaceAfter=7,
        uppercase=True,
    )
    body_style = ParagraphStyle(
        "AnalyticsBody",
        parent=styles["BodyText"],
        textColor=MUTED,
        fontSize=8,
        leading=12,
    )
    table_header = ParagraphStyle(
        "AnalyticsTableHeader",
        parent=body_style,
        textColor=colors.white,
        fontName="Helvetica-Bold",
    )

    generated_at = datetime.now().strftime("%B %-d, %Y at %-I:%M %p")
    included_pages = ", ".join(PAGE_LABELS[page] for page in pages)
    cover = Table(
        [[
            Paragraph("RSS ANALYTICS", section_label),
            Paragraph("Decision-ready report", ParagraphStyle(
                "CoverEyebrowRight",
                parent=section_label,
                textColor=GOLD,
                alignment=TA_RIGHT,
            )),
        ], [
            Paragraph("Analytics Report", title_style),
            "",
        ], [
            Paragraph(f"Generated {generated_at}", cover_subtitle),
            "",
        ], [
            Paragraph(f"Included pages: {included_pages}", cover_subtitle),
            "",
        ]],
        colWidths=[4.9 * inch, 4.9 * inch],
        rowHeights=[0.35 * inch, 0.75 * inch, 0.35 * inch, 0.55 * inch],
    )
    cover.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), NAVY),
        ("SPAN", (0, 1), (-1, 1)),
        ("SPAN", (0, 2), (-1, 2)),
        ("SPAN", (0, 3), (-1, 3)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 18),
        ("RIGHTPADDING", (0, 0), (-1, -1), 18),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story = [cover, Spacer(1, 18)]

    filter_labels = {
        "date_from": "Current Period From",
        "date_to": "Current Period To",
        "comparison_date_from": "Previous Period From",
        "comparison_date_to": "Previous Period To",
        "location": "Location",
        "locations": "Location",
        "lead_source": "Lead Source",
        "lead_sources": "Lead Source",
        "staff": "Staff",
    }
    filter_rows = []
    for key, value in sorted(params.items()):
        if key not in filter_labels:
            continue
        display_value = ", ".join(str(item) for item in value) if isinstance(value, list) else value
        if display_value not in (None, ""):
            filter_rows.append([
                Paragraph(filter_labels[key], body_style),
                Paragraph(str(display_value), ParagraphStyle(
                    f"FilterValue{key}",
                    parent=body_style,
                    textColor=NAVY_DARK,
                    fontName="Helvetica-Bold",
                )),
            ])
    if filter_rows:
        story.extend([
            Paragraph("REPORT CONTEXT", section_label),
            _styled_pdf_table(filter_rows, [1.8 * inch, 7.8 * inch], header=False),
        ])

    for index, page in enumerate(pages):
        if len(pages) > 1 or index:
            story.append(PageBreak())
        else:
            story.append(Spacer(1, 18))
        story.extend([
            Paragraph(PAGE_LABELS[page], page_title),
            Paragraph(
                "Current selections and authorized data captured when this report was generated.",
                body_style,
            ),
            Spacer(1, 8),
        ])
        if page == "overview":
            story.extend(_overview_pdf_content(data, section_label, body_style))
            continue
        story.append(Paragraph("PERFORMANCE DETAILS", section_label))
        rows = _rows(data, page)
        column_count = max(len(row) for row in rows)
        headers = rows[0]
        normalized_rows = [
            [
                Paragraph(
                    escape(_pdf_display_value(page, row_index, column_index, value, row, headers)),
                    table_header if row_index == 0 else body_style,
                )
                for column_index, value in enumerate(row)
            ] + [Paragraph("", body_style)] * (column_count - len(row))
            for row_index, row in enumerate(rows)
        ]
        usable_width = 9.96 * inch
        story.append(_styled_pdf_table(
            normalized_rows,
            [usable_width / column_count] * column_count,
            header=True,
        ))
        story.append(Spacer(1, 8))
        story.append(KeepTogether([
            Paragraph("How to read this page", section_label),
            Paragraph(
                "Values follow the same analytics business rules and role permissions as the RSS workspace. "
                "Interactive controls, hover states, and drill-through actions are intentionally omitted from the PDF.",
                body_style,
            ),
        ]))

    document.build(
        story,
        onFirstPage=_pdf_header_footer,
        onLaterPages=_pdf_header_footer,
    )
    return stream.getvalue()


def _overview_pdf_content(data, section_label, body_style):
    labels = [
        ("scheduled", "Booked"),
        ("toured", "Toured"),
        ("no_show", "No Show"),
        ("enrolled", "Enrolled"),
        ("churned", "Churned"),
    ]
    current = data.get("volumeCounts") or data.get("counts", {})
    previous = data.get("previousVolumeCounts") or data.get("previousCounts", {})
    deltas = data.get("volumeDeltas") or data.get("deltas", {})
    rate_labels = [
        ("toured", "Tour Rate", False),
        ("no_show", "No-show Rate", True),
        ("close", "Close Rate", False),
        ("conversion", "Conversion Rate", False),
    ]
    rates = data.get("rates", {})
    rate_deltas = data.get("rateDeltas", {})

    flowables = [Paragraph("VOLUME AND TREND ANALYSIS", section_label)]
    flowables.append(_pdf_card_row([
        (
            label,
            f"{current.get(key, 0):,}",
            _volume_delta_text(deltas.get(key), current.get(key, 0), previous.get(key, 0)),
        )
        for key, label in labels
    ], body_style))
    flowables.extend([Spacer(1, 12), Paragraph("CONVERSION AND COHORT ANALYSIS", section_label)])
    rate_cards = [
        (
            label,
            "-" if rates.get(key) is None else f"{rates[key]:.1f}%",
            _rate_delta_text(rate_deltas.get(key), inverse),
        )
        for key, label, inverse in rate_labels
    ]
    average_days = data.get("averageDaysToEnroll")
    rate_cards.append((
        "Average Days to Enroll",
        "-" if average_days is None else f"{average_days:.1f}",
        "Scheduled tour to enrollment",
    ))
    flowables.append(_pdf_card_row(rate_cards, body_style))

    pending = data.get("pendingCounts", {})
    rankings = data.get("rankings", {})
    financial = data.get("financialSummary", {})
    insight_rows = [
        ["Operational signal", "Value", "Interpretation"],
        [
            "Past tours awaiting an outcome",
            str(pending.get("pendingTourOutcome", 0)),
            "Review tours that have not reached Toured or No Show.",
        ],
        [
            "Toured families awaiting a final outcome",
            str(pending.get("pendingEnrollmentOutcome", 0)),
            "Review toured families that have not reached Enrolled or Churned.",
        ],
    ]
    top_entities = [
        ("Top Location", rankings.get("locations", [])),
        ("Top Lead Source", rankings.get("leadSources", rankings.get("lead_sources", []))),
        ("Top Staff", rankings.get("staff", [])),
    ]
    for label, rows in top_entities:
        if rows:
            item = rows[0]
            insight_rows.append([
                label,
                str(item.get("name", item.get("label", "-"))),
                _ranking_summary(item),
            ])
    if any(float(financial.get(key, 0) or 0) for key in ("revenue", "cost", "margin")):
        insight_rows.append([
            "Contribution Margin",
            f"${float(financial.get('margin', 0) or 0):,.0f}",
            f"Revenue ${float(financial.get('revenue', 0) or 0):,.0f} less costs ${float(financial.get('cost', 0) or 0):,.0f}.",
        ])

    header_style = ParagraphStyle(
        "OverviewInsightHeader",
        parent=body_style,
        textColor=colors.white,
        fontName="Helvetica-Bold",
    )
    normalized = [
        [
            Paragraph(escape(str(value)), header_style if row_index == 0 else body_style)
            for value in row
        ]
        for row_index, row in enumerate(insight_rows)
    ]
    flowables.extend([
        Spacer(1, 12),
        KeepTogether([
            Paragraph("PERFORMANCE INSIGHTS", section_label),
            _styled_pdf_table(normalized, [2.7 * inch, 2.15 * inch, 5.1 * inch], header=True),
        ]),
    ])
    return flowables


def _pdf_card_row(cards, body_style):
    card_style = ParagraphStyle(
        "AnalyticsCard",
        parent=body_style,
        textColor=NAVY_DARK,
        fontSize=8,
        leading=13,
        alignment=TA_CENTER,
    )
    cells = []
    for label, value, detail in cards:
        cells.append(Paragraph(
            f"<font color='#65758D' size='7'><b>{escape(label)}</b></font><br/>"
            f"<font color='#071B40' size='16'><b>{escape(str(value))}</b></font><br/>"
            f"<font color='#65758D' size='6'>{escape(detail or 'No comparison available')}</font>",
            card_style,
        ))
    table = Table([cells], colWidths=[9.96 * inch / len(cells)] * len(cells), hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), SOFT_BLUE),
        ("BOX", (0, 0), (-1, -1), 0.75, GRID),
        ("INNERGRID", (0, 0), (-1, -1), 0.75, colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def _volume_delta_text(delta_value, current, previous):
    if isinstance(delta_value, dict):
        difference = delta_value.get("difference")
        percent = delta_value.get("percent")
    else:
        difference = current - previous
        percent = round((difference / previous) * 100) if previous else None
    if difference is None:
        return "No comparison available"
    sign = "+" if difference > 0 else ""
    percent_text = "" if percent is None else f" ({sign}{percent:.0f}%)"
    return f"Previous {previous:,} - {sign}{difference:,}{percent_text}"


def _rate_delta_text(delta_value, inverse=False):
    if delta_value is None:
        return "No comparison available"
    direction = "improvement" if (delta_value < 0 if inverse else delta_value > 0) else "needs attention"
    sign = "+" if delta_value > 0 else ""
    return f"{sign}{delta_value:.1f} pts - {direction}"


def _ranking_summary(item):
    value = item.get("value")
    if value is None:
        return "Top authorized result for the selected ranking."
    metric = item.get("metric", "selected metric").replace("_", " ")
    return f"{value:.1f} for {metric}."


def _pdf_display_value(page, row_index, column_index, value, row, headers):
    if value is None:
        return "—"
    if row_index == 0 or value == "":
        return str(value)
    header = str(headers[column_index] if column_index < len(headers) else "").lower()
    row_label = str(row[0] if row else "").lower()
    if isinstance(value, (int, float)):
        if page == "cost-margin":
            is_currency = (
                column_index == 1 and row_label in {"revenue", "costs", "contribution margin"}
            ) or any(label in header for label in ("revenue", "costs", "contribution margin")) or (
                row_index >= 5 and column_index in {1, 2, 3}
            )
            if is_currency:
                return f"${value:,.0f}"
            if "margin %" in header or (row_index >= 5 and column_index == 4):
                return f"{value:.1f}%"
        if page == "cohort" and column_index == 1 and "days" not in row_label:
            return f"{value:.1f}%"
        if page in {"locations", "leadSources", "staff"} and "conversion" in header:
            return f"{value:.1f}%"
        if isinstance(value, float):
            return f"{value:,.1f}"
        return f"{value:,}"
    return str(value)


def _styled_pdf_table(rows, widths, header=True):
    table = Table(rows, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    style = [
        ("GRID", (0, 0), (-1, -1), 0.35, GRID),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("ROWBACKGROUNDS", (0, 1 if header else 0), (-1, -1), [colors.white, colors.HexColor("#F7FAFD")]),
    ]
    if header:
        style.extend([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ])
    table.setStyle(TableStyle(style))
    return table


def _pdf_header_footer(canvas, document):
    canvas.saveState()
    width, height = landscape(letter)
    canvas.setFillColor(NAVY)
    canvas.rect(0, height - 0.42 * inch, width, 0.42 * inch, fill=1, stroke=0)
    canvas.setFillColor(GOLD)
    canvas.setFont("Helvetica-Bold", 8)
    canvas.drawString(0.52 * inch, height - 0.27 * inch, "RSS ANALYTICS")
    canvas.setFillColor(MUTED)
    canvas.setFont("Helvetica", 7)
    canvas.drawRightString(width - 0.52 * inch, 0.28 * inch, f"Page {document.page}")
    canvas.setStrokeColor(GRID)
    canvas.line(0.52 * inch, 0.43 * inch, width - 0.52 * inch, 0.43 * inch)
    canvas.restoreState()
